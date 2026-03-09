#!/usr/bin/env python3
from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook

from pipeline_state import resolve_workbook_path

SYMBOL_KEYS = (
    "symbol",
    "ticker",
    "code",
    "代码",
    "证券代码",
    "股票代码",
    "标的代码",
)
QTY_KEYS = (
    "quantity",
    "qty",
    "shares",
    "持仓数量",
    "持仓",
    "数量",
    "持股",
    "持股数",
    "总数量",
    "可用数量",
)


def _norm_col(name: object) -> str:
    return str(name or "").strip().lower().replace(" ", "")


def _find_col(columns: Iterable[object], candidates: tuple[str, ...]) -> object | None:
    norm_map = {c: _norm_col(c) for c in columns}
    candidate_norm = {_norm_col(c) for c in candidates}
    for col, norm in norm_map.items():
        if norm in candidate_norm:
            return col
    return None


def _to_float(v: object) -> float:
    if pd.isna(v):
        return 0.0
    s = str(v).strip().replace(",", "")
    if not s:
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def _normalize_symbol(raw: object) -> str:
    s = str(raw or "").strip().upper()
    if not s:
        return ""
    s = s.replace(" ", "")
    if s in {"CASH", "USD", "USDT", "现金"}:
        return "CASH"
    if s in {"GOLD", "黄金", "AU"}:
        return "GOLD.CN"
    if s.endswith(".US") or s.endswith(".CN"):
        return s
    if s.isalpha() and 1 <= len(s) <= 6:
        return f"{s}.US"
    return s


def extract_positions_from_file(path: Path) -> dict[str, float]:
    out: dict[str, float] = defaultdict(float)
    xls = pd.ExcelFile(path)
    for sheet in xls.sheet_names:
        df = pd.read_excel(path, sheet_name=sheet)
        if df.empty:
            continue
        sym_col = _find_col(df.columns, SYMBOL_KEYS)
        qty_col = _find_col(df.columns, QTY_KEYS)
        if sym_col is None or qty_col is None:
            continue
        for _, row in df.iterrows():
            sym = _normalize_symbol(row.get(sym_col))
            qty = _to_float(row.get(qty_col))
            if not sym:
                continue
            out[sym] += qty
    return dict(out)


def apply_to_assets(
    *,
    positions: dict[str, float],
    keep_symbols: set[str],
    zero_missing: bool,
) -> dict[str, object]:
    asset_xlsx = resolve_workbook_path()
    wb = load_workbook(asset_xlsx)
    ws = wb["Holdings"]

    headers = {str(ws.cell(1, c).value).strip().lower(): c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    ts_col = headers.get("timestamp")
    sym_col = headers.get("symbol")
    qty_col = headers.get("quantity")
    if not sym_col or not qty_col:
        raise RuntimeError("Holdings sheet missing required columns: symbol/quantity")

    changed: list[str] = []
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    seen = set()

    for r in range(2, ws.max_row + 1):
        sym = _normalize_symbol(ws.cell(r, sym_col).value)
        if not sym:
            continue
        seen.add(sym)
        old = _to_float(ws.cell(r, qty_col).value)
        if sym in positions:
            new = positions[sym]
        elif zero_missing and sym not in keep_symbols and sym not in {"CASH"}:
            new = 0.0
        else:
            new = old
        if abs(new - old) > 1e-12:
            ws.cell(r, qty_col).value = new
            changed.append(f"{sym}: {old} -> {new}")
        if ts_col:
            ws.cell(r, ts_col).value = now

    # append new symbols not in template
    for sym, qty in positions.items():
        if sym in seen:
            continue
        if qty <= 0:
            continue
        nr = ws.max_row + 1
        if ts_col:
            ws.cell(nr, ts_col).value = now
        ws.cell(nr, sym_col).value = sym
        ws.cell(nr, qty_col).value = qty
        changed.append(f"{sym}: new -> {qty}")

    wb.save(asset_xlsx)
    return {"changed": changed, "changed_count": len(changed), "timestamp": now}


def main() -> int:
    ap = argparse.ArgumentParser(description="Import broker position exports into Asset-Management/assets.xlsx Holdings")
    ap.add_argument("--source", action="append", required=True, help="Input xlsx file path (repeatable)")
    ap.add_argument("--keep-symbol", action="append", default=["GOLD.CN"], help="Symbols to keep unchanged if missing")
    ap.add_argument("--no-zero-missing", action="store_true", help="Do not zero out missing non-cash symbols")
    args = ap.parse_args()

    all_pos: dict[str, float] = defaultdict(float)
    for src in args.source:
        path = Path(src).expanduser()
        if not path.exists():
            raise FileNotFoundError(f"source not found: {path}")
        extracted = extract_positions_from_file(path)
        for k, v in extracted.items():
            all_pos[k] += v

    result = apply_to_assets(
        positions=dict(all_pos),
        keep_symbols={_normalize_symbol(s) for s in args.keep_symbol},
        zero_missing=not args.no_zero_missing,
    )
    print(f"updated {result['changed_count']} row(s) at {result['timestamp']}")
    for line in result["changed"][:30]:
        print(f"- {line}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
